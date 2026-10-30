import copy
import csv
import io
import json
import os
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parents[1]))
from core import ExtractionError, apply_page_edits, clock, extract_pdf, render_page, validate_result
from exports import build_exports


@pytest.fixture(scope="session")
def pdf_data():
    path = os.environ.get("JADUAL_TEST_PDF")
    if not path or not Path(path).is_file():
        pytest.skip("Set JADUAL_TEST_PDF kepada PDF rujukan untuk ujian integrasi.")
    return Path(path).read_bytes()


@pytest.fixture(scope="session")
def result(pdf_data):
    return extract_pdf(pdf_data, "rujukan.pdf")


def test_reference_counts_and_two_reader_checks(result):
    check = validate_result(result)
    assert (check["teachers"], check["lessons"], check["blocks"], check["supported_pages"]) == (75, 2031, 1214, 75)
    assert len(check["subject_checks"]) == 174
    assert check["blockers"] == check["reviews"] == 0
    assert sum(p["text_checks"] for p in result["pages"]) == 1214
    assert not any(p["text_failures"] for p in result["pages"])


def test_matches_independently_saved_reference_rows(result):
    reference = os.environ.get("JADUAL_TEST_REFERENCE")
    if not reference:
        pytest.skip("Set JADUAL_TEST_REFERENCE kepada tables.json ekstrak rujukan.")
    tables = json.loads(Path(reference).read_text(encoding="utf-8"))
    expected = [dict(zip(tables["Jadual"][0], r)) for r in tables["Jadual"][1:]]
    for a, b in zip(result["rows"], expected, strict=True):
        for key, value in b.items():
            if key in ("masa_mula", "masa_tamat"):
                value = clock(value)
            assert a[key] == value, (a["jadual_id"], key, a[key], value)


@pytest.mark.parametrize("data", [b"", b"not a PDF", b"x"*(20*1024*1024+1), b"%PDF-invalid"])
def test_rejects_invalid_input(data):
    with pytest.raises(ExtractionError):
        extract_pdf(data)


def test_repeated_teacher_page_is_not_silently_exported(pdf_data):
    from pypdf import PdfReader, PdfWriter
    reader, writer = PdfReader(io.BytesIO(pdf_data)), PdfWriter()
    writer.add_page(reader.pages[0]); writer.add_page(reader.pages[0])
    out = io.BytesIO(); writer.write(out)
    r = extract_pdf(out.getvalue())
    assert validate_result(r)["blockers"] > 0
    with pytest.raises(ValueError, match="disekat"):
        build_exports(r, confirmed=True)


def test_scaled_vector_layout_is_detected_from_geometry(pdf_data):
    from pypdf import PdfReader, PdfWriter
    reader, writer = PdfReader(io.BytesIO(pdf_data)), PdfWriter()
    writer.add_page(reader.pages[0]); writer.pages[0].scale_by(.8)
    out = io.BytesIO(); writer.write(out)
    r = extract_pdf(out.getvalue())
    assert validate_result(r)["blockers"] == 0
    assert len(r["rows"]) == 14


def test_image_pdf_explains_ocr_requirement():
    from PIL import Image
    out = io.BytesIO(); Image.new("RGB", (800, 500), "white").save(out, "PDF")
    r = extract_pdf(out.getvalue())
    assert validate_result(r)["blockers"] > 0
    assert any("OCR" in i["mesej"] for i in r["issues"])


def test_encrypted_pdf_reports_clear_error(pdf_data):
    from pypdf import PdfReader, PdfWriter
    reader, writer = PdfReader(io.BytesIO(pdf_data)), PdfWriter()
    writer.add_page(reader.pages[0]); writer.encrypt("secret")
    out = io.BytesIO(); writer.write(out)
    with pytest.raises(ExtractionError, match="kata laluan"):
        extract_pdf(out.getvalue())


@pytest.mark.parametrize("change", [
    {"masa_tamat": "08:00"}, {"masa_mula": "24:00"}, {"kelas_label": ""},
    {"hari_no": 1.5}, {"waktu": 0}, {"nama_guru": "Nama tidak sepadan"},
])
def test_invalid_row_blocks_export(result, change):
    r = copy.deepcopy(result); r["rows"][0].update(change)
    assert validate_result(r)["blockers"] > 0
    with pytest.raises(ValueError):
        build_exports(r, confirmed=True)


def test_duplicate_slot_and_real_time_overlap(result):
    r = copy.deepcopy(result); r["rows"].append(copy.deepcopy(r["rows"][0]))
    assert validate_result(r)["blockers"] > 0
    r = copy.deepcopy(result)
    # Waktu 4 tamat 09:30; waktu 5 dimulakan 09:20. Kedua-duanya bertindih.
    r["rows"][1]["masa_mula"] = "09:20"
    assert any(i["kod"] == "PERTINDIHAN" for i in validate_result(r)["issues"])


def test_overlap_reports_exact_rows_times_classes_and_intersection(result):
    r = copy.deepcopy(result)
    r["rows"][1]["masa_mula"] = "09:20"
    issue = next(i for i in validate_result(r)["issues"] if i["kod"] == "PERTINDIHAN")
    assert (issue["baris_a"], issue["baris_b"]) == (1, 2)
    assert issue["masa_a"] == "09:00-09:30"
    assert issue["masa_b"] == "09:20-10:00"
    assert issue["julat_bertindih"] == "09:20-09:30"
    assert issue["rekod_a"] == r["rows"][0]["rekod_id"]
    assert "3H" in issue["mesej"]


def test_long_overlap_is_found_even_with_intervening_short_slot(result):
    r = copy.deepcopy(result)
    r["rows"] = r["rows"][:3]
    r["rows"][0]["masa_tamat"] = "12:00"
    issues = [i for i in validate_result(r)["issues"] if i["kod"] == "PERTINDIHAN"]
    assert any(i["baris_a"] == 1 and i["baris_b"] == 3 for i in issues)


def test_manual_class_and_name_correction_is_tracked_without_changing_source(result):
    rows = copy.deepcopy([r for r in result["rows"] if r["halaman_pdf"] == 1])
    rows[0]["kelas_label"] = "3Z"
    updated = apply_page_edits(result, 1, rows, "CIKGU ZUNAINUN")
    assert updated["rows"][0]["kelas_label"] == "3Z"
    assert updated["rows"][0]["nama_guru"] == "CIKGU ZUNAINUN"
    assert result["rows"][0]["kelas_label"] == "3H"
    assert updated["blocks"][0]["kelas_asal"] == result["blocks"][0]["kelas_asal"]
    assert len(updated["edits"]) == 14
    assert validate_result(updated)["blockers"] == 0


def test_invalid_edit_does_not_mutate_original(result):
    rows = copy.deepcopy([r for r in result["rows"] if r["halaman_pdf"] == 1])
    rows[0]["waktu"] = 0
    with pytest.raises(ValueError):
        apply_page_edits(result, 1, rows, result["teachers"][0]["nama_guru"])
    assert result["rows"][0]["waktu"] == 4


def test_review_needs_explicit_reason_and_confirmation(result):
    r = copy.deepcopy(result); r["rows"][0]["subjek_label"] = "SUBJEK_BAHARU"
    assert validate_result(r)["reviews"] > 0
    with pytest.raises(ValueError, match="Sahkan"):
        build_exports(r)
    with pytest.raises(ValueError, match="10 aksara"):
        build_exports(r, confirmed=True)
    out = build_exports(r, confirmed=True, review_reason="Label telah disemak secara manual pada PDF.")
    assert json.loads(out["json"])["semakan"]["status"] == "DISAHKAN_PENGGUNA"


def test_exports_are_complete_typed_and_literal(result):
    import openpyxl
    exports = build_exports(result, confirmed=True)
    workbook = openpyxl.load_workbook(io.BytesIO(exports["xlsx"]), read_only=True, data_only=False)
    assert workbook.sheetnames[:2] == ["Jadual", "Guru"]
    sheet = workbook["Jadual"]
    assert sheet.max_row == 2032
    assert workbook["Guru"].max_row == 76
    assert sheet["G2"].value.strftime("%H:%M") == "09:00"
    assert sheet["J2"].value == "3H"
    assert not any(c.data_type in ("f", "e") for ws in workbook for row in ws for c in row)
    workbook.close()
    data = json.loads(exports["json"])
    assert len(data["jadual"]) == 2031 and len(data["guru"]) == 75
    rows = list(csv.DictReader(io.StringIO(exports["csv"].decode("utf-8-sig"))))
    assert len(rows) == 2031 and rows[0]["masa_mula"] == "09:00"
    with zipfile.ZipFile(io.BytesIO(exports["zip"])) as archive:
        assert archive.testzip() is None
        assert "jadual.xlsx" in archive.namelist()
        assert "laporan_semakan.json" in archive.namelist()
    # Malicious-looking text is literal in Excel and quoted for CSV imports.
    altered = copy.deepcopy(result); altered["rows"][0]["kelas_label"] = '=HYPERLINK("x")'
    out = build_exports(altered, confirmed=True)
    wb = openpyxl.load_workbook(io.BytesIO(out["xlsx"]), read_only=True, data_only=False)
    assert wb["Jadual"]["J2"].data_type == "s"; wb.close()
    assert list(csv.DictReader(io.StringIO(out["csv"].decode("utf-8-sig"))))[0]["kelas_label"].startswith("'=")


def test_source_image_rendering(pdf_data, result):
    from PIL import Image
    png = render_page(pdf_data, 51, result["blocks"], "G051-B02")
    image = Image.open(io.BytesIO(png))
    assert image.width > 1000 and image.height > 800
